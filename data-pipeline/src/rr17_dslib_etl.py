"""
RR-17 part 1 — ETL for the Ghent DSLIB v3.4 corpus (Batselier & Vanhoucke 2015, IJPM 33(3)).

Why DSLIB: it is the only open corpus we found with (a) BUILDING projects (own domain), (b) a FROZEN baseline
schedule, and (c) periodic tracking snapshots (TP1..TPn) with actual start / actual duration per activity.
JPF (RR-15/16) has neither buildings nor snapshots; its planned dates move with the plan.

Input : DSLIB 3.4/Excel/*.xlsx  (one workbook per project; sheets 'Baseline Schedule', 'Project Control - TP1', 'TP2'..)
        DSLIB 3.4/DSLIB_Analysis_Sheet.xlsx (sector, authenticity flags G/Y/O)
Output: outputs/dslib_labeled_tasks.csv   — one row per LEAF activity, same 11 plan-time features as etl.py
        outputs/dslib_snapshots.csv       — one row per (project, TP, activity): status at the TP status date
        outputs/dslib_etl_summary.json

Feature derivation mirrors data-pipeline/src/etl.py (names, BFS cap 5000):
  planned_duration_days = baseline end − baseline start (calendar days)
  total_float_hr / free_float_hr = CPM forward/backward pass on the baseline network in WORKING days × 8h
     (DSLIB exports no float column; lags parsed from '12FS-2w 4d' — w=5 wd, d=1 wd, h=1/8 wd; ProTrack default agenda)
  rel_position, proj_n_tasks, proj_span_days, n_pred, n_succ, upstream_cnt, downstream_cnt — as etl.py
  task_type = None (DSLIB has no P6 task type; the pipeline imputes "NA", as for JPF)
Labels (from the LAST tracking period, finished activities only — PC=1 or status 'Finished'):
  actual_end ≈ actual_start + actual_duration in working days (Mon–Fri; holidays ignored — approximation, documented)
  delay_days   = actual_end − baseline_end (calendar days)         is_late   = delay_days > 0   (frozen-baseline label)
  start_delay  = actual_start − baseline_start (calendar days)     start_late = start_delay > 0
  grew         = actual_duration > baseline_duration (working days) (the activity's OWN overrun, plan-drift removed)
  own_delay    = delay_days − start_delay > 0                       (≈ grew, calendar version)
Only CONSTRUCTION sectors are kept (own domain + civil); IT/Engineering/Event/Mobility/Education are dropped.
Run:  python rr17_dslib_etl.py "<DSLIB 3.4 dir>" [out_dir]
"""
from __future__ import annotations
import re, json, sys, warnings
from collections import defaultdict, deque
from datetime import datetime, timedelta
from pathlib import Path
import numpy as np, pandas as pd
warnings.filterwarnings("ignore")
import openpyxl

ROOT = Path(sys.argv[1]); OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "outputs"); OUT.mkdir(parents=True, exist_ok=True)
HOURS_PER_DAY, BFS_CAP = 8.0, 5000

# ---------- analysis sheet ----------
meta = {}
for r in list(openpyxl.load_workbook(ROOT / "DSLIB_Analysis_Sheet.xlsx", read_only=True, data_only=True)["DSLIB"].iter_rows(values_only=True))[3:]:
    if r[0] and isinstance(r[0], str) and re.match(r"C20\d\d-\d\d", r[0]):
        meta[r[0]] = dict(name=r[1], sector=r[3], auth_project=r[8], auth_tracking=r[9])

def sector_group(s):
    s = (s or "").lower()
    if not s.startswith("construction"): return None
    return "buildings" if "building" in s else ("civil" if "civil" in s else "industrial")

# ---------- parsers ----------
def todt(v):
    if v is None or isinstance(v, datetime): return v
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y %H:%M", "%m/%d/%Y"):
        try: return datetime.strptime(str(v).strip(), fmt)
        except ValueError: pass
    return None

UNIT = {"w": 5.0, "d": 1.0, "h": 1 / HOURS_PER_DAY}
def span_wd(s):
    """'2w 4d 6h' → working days; plain number → days."""
    if s is None: return None
    if isinstance(s, (int, float)): return float(s)
    s = str(s).strip()
    if re.fullmatch(r"\d+(\.\d+)?", s): return float(s)
    tot, ok = 0.0, False
    for n, u in re.findall(r"(\d+(?:\.\d+)?)\s*([wdh])", s):
        tot += float(n) * UNIT[u]; ok = True
    return tot if ok else None

PRED = re.compile(r"^(\d+)(FS|SS|FF|SF)(?:([+-])\s*(.*))?$")
def parse_preds(cell):
    out = []
    for tok in str(cell).split(";") if cell else []:
        m = PRED.match(tok.strip())
        if not m: continue
        lag = span_wd(m.group(4)) or 0.0
        out.append((int(m.group(1)), m.group(2), lag if m.group(3) != "-" else -lag))
    return out

def add_wd(start, wdays, weekend=(5, 6)):
    if wdays is None or start is None: return None
    d = start; n = int(round(wdays))
    if n <= 0: return d
    n -= 1
    while n > 0:
        d += timedelta(days=1)
        if d.weekday() not in weekend: n -= 1
    return d

def cpm_floats(ids, dur, preds):
    """Total/free float in working days. preds[j] = [(i, type, lag)] meaning i → j. Works for DAGs; cycles → NaN."""
    succ = defaultdict(list)
    for j, lst in preds.items():
        for i, t, lag in lst:
            if i in ids: succ[i].append((j, t, lag))
    indeg = {t: 0 for t in ids}
    for i in succ:
        for j, _, _ in succ[i]: indeg[j] += 1
    order, q = [], deque([t for t in ids if indeg[t] == 0])
    while q:
        u = q.popleft(); order.append(u)
        for v, _, _ in succ[u]:
            indeg[v] -= 1
            if indeg[v] == 0: q.append(v)
    if len(order) != len(ids): return {t: (np.nan, np.nan) for t in ids}
    ES = {t: 0.0 for t in ids}
    for j in order:
        for i, t, lag in preds.get(j, []):
            if i not in ids: continue
            d_i, d_j = dur[i], dur[j]
            cand = {"FS": ES[i] + d_i + lag, "SS": ES[i] + lag, "FF": ES[i] + d_i + lag - d_j, "SF": ES[i] + lag - d_j}[t]
            ES[j] = max(ES[j], cand)
    EF = {t: ES[t] + dur[t] for t in ids}
    horizon = max(EF.values()) if EF else 0.0
    LF = {t: horizon for t in ids}
    for i in reversed(order):
        for j, t, lag in succ[i]:
            d_i, d_j = dur[i], dur[j]
            LS_j = LF[j] - d_j
            cand = {"FS": LS_j - lag, "SS": LS_j - lag + d_i, "FF": LF[j] - lag, "SF": LF[j] - lag + d_i}[t]
            LF[i] = min(LF[i], cand)
    TF = {t: LF[t] - EF[t] for t in ids}
    FF = {}
    for i in ids:
        if not succ[i]: FF[i] = TF[i]; continue
        slack = []
        for j, t, lag in succ[i]:
            d_i, d_j = dur[i], dur[j]
            slack.append({"FS": ES[j] - (EF[i] + lag), "SS": ES[j] - (ES[i] + lag),
                          "FF": (ES[j] + d_j) - (EF[i] + lag), "SF": (ES[j] + d_j) - (ES[i] + lag)}[t])
        FF[i] = max(0.0, min(slack))
    return {t: (TF[t], FF[t]) for t in ids}

def graph_stats(ids, edges):
    succ, pred = defaultdict(list), defaultdict(list)
    for a, b in edges:
        if a in ids and b in ids: succ[a].append(b); pred[b].append(a)
    def reach(start, adj, cap=BFS_CAP):
        seen, q = {start}, deque([start])
        while q and len(seen) < cap:
            for n in adj[q.popleft()]:
                if n not in seen: seen.add(n); q.append(n)
        return len(seen) - 1
    return {t: dict(n_pred=len(pred[t]), n_succ=len(succ[t]), upstream_cnt=reach(t, pred), downstream_cnt=reach(t, succ)) for t in ids}

# ---------- main loop ----------
rows, snaps, proj_summ = [], [], []
cycles = 0; derived_end = [0]
for f in sorted((ROOT / "Excel").glob("*.xlsx")):
    code = f.name.split(" ")[0]; m = meta.get(code, {})
    grp = sector_group(m.get("sector"))
    if grp is None: continue
    w = openpyxl.load_workbook(f, read_only=True, data_only=True)
    bs = list(w["Baseline Schedule"].iter_rows(values_only=True))
    recs = [r for r in bs[2:] if r and r[0] not in (None, 0)]
    wbs = {r[0]: str(r[2]) if r[2] is not None else "" for r in recs}
    allw = set(v for v in wbs.values() if v)
    leaf = {k: bool(v) and not any(o != v and o.startswith(v + ".") for o in allw) for k, v in wbs.items()}
    if not allw: leaf = {k: True for k in wbs}  # 8 newer workbooks export no WBS column → every non-summary row is a leaf
    ids = {r[0] for r in recs if leaf[r[0]]}
    base = {r[0]: dict(name=r[1], b_start=todt(r[5]), b_end=todt(r[6]), b_dur=span_wd(r[7]), preds=parse_preds(r[3])) for r in recs}
    for b in base.values():  # same workbooks omit Baseline End → derive it from start + working-day duration
        if b["b_end"] is None and b["b_start"] is not None and b["b_dur"] is not None: b["b_end"] = add_wd(b["b_start"], b["b_dur"]); derived_end[0] += 1
    preds = {k: [(i, t, lag) for i, t, lag in v["preds"] if i in ids] for k, v in base.items() if k in ids}
    dur = {k: (base[k]["b_dur"] or 0.0) for k in ids}
    fl = cpm_floats(ids, dur, preds)
    if any(np.isnan(v[0]) for v in fl.values()): cycles += 1
    edges = [(i, j) for j, lst in preds.items() for i, _, _ in lst]
    g = graph_stats(ids, edges)
    starts = [base[k]["b_start"] for k in ids if base[k]["b_start"]]; ends = [base[k]["b_end"] for k in ids if base[k]["b_end"]]
    p0, p1 = (min(starts) if starts else None), (max(ends) if ends else None)
    span = max((p1 - p0).days, 1) if p0 and p1 else np.nan

    tps = [s for s in w.sheetnames if s.startswith("Project Control - TP") or re.fullmatch(r"TP\d+", s)]
    last = {}
    for k, sn in enumerate(tps, 1):
        tr = list(w[sn].iter_rows(values_only=True)); sd = todt(tr[0][2])
        for r in tr[4:]:
            if not r or r[0] not in ids: continue
            a_start, a_dur, pc, st = todt(r[11]), span_wd(r[12]), r[20], str(r[21] or "").strip().lower()
            done = (pc == 1) or (st == "finished")
            status = "finished" if done else ("in_progress" if a_start is not None else "not_started")
            snaps.append(dict(project=code, tp=k, status_date=sd, task_id=r[0], status=status, actual_start=a_start,
                              actual_duration_wd=a_dur if status != "not_started" else np.nan,
                              pct_complete=pc if isinstance(pc, (int, float)) else np.nan))
            if k == len(tps): last[r[0]] = (a_start, a_dur, done)
    for k in ids:
        b = base[k]; a_start, a_dur, done = last.get(k, (None, None, False))
        labeled = done and a_start is not None and a_dur is not None and b["b_start"] and b["b_end"]
        a_end = add_wd(a_start, a_dur) if labeled else None
        delay = (a_end.date() - b["b_end"].date()).days if labeled else np.nan
        sdel = (a_start.date() - b["b_start"].date()).days if labeled else np.nan
        rows.append(dict(
            project=code, source="dslib", sector=m.get("sector"), sector_group=grp, auth_tracking=m.get("auth_tracking"),
            task_id=k, task_name=b["name"], task_type=None, baseline_start=b["b_start"], baseline_end=b["b_end"],
            planned_duration_days=((b["b_end"] - b["b_start"]).days if b["b_start"] and b["b_end"] else np.nan),
            planned_duration_wd=b["b_dur"],
            total_float_hr=fl[k][0] * HOURS_PER_DAY, free_float_hr=fl[k][1] * HOURS_PER_DAY,
            rel_position=(((b["b_start"] - p0).days / span) if b["b_start"] and p0 and span == span else np.nan),
            proj_n_tasks=len(ids), proj_span_days=span, n_tp=len(tps), **g[k],
            actual_start=a_start if labeled else None, actual_duration_wd=a_dur if labeled else np.nan, actual_end=a_end,
            delay_days=delay, start_delay_days=sdel,
            is_late=(float(delay > 0) if labeled else np.nan), is_late_7=(float(delay > 7) if labeled else np.nan),
            start_late=(float(sdel > 0) if labeled else np.nan),
            grew=(float(a_dur > (b["b_dur"] or 0)) if labeled else np.nan),
            own_delay=(float((delay - sdel) > 0) if labeled else np.nan)))
    proj_summ.append(dict(project=code, name=m.get("name"), sector_group=grp, n_leaf=len(ids), n_tp=len(tps),
                          labeled=sum(1 for k in ids if last.get(k, (None, None, False))[2] and last[k][0] is not None)))

df = pd.DataFrame(rows); sn = pd.DataFrame(snaps)
df.to_csv(OUT / "dslib_labeled_tasks.csv", index=False); sn.to_csv(OUT / "dslib_snapshots.csv", index=False)
lab = df[df["is_late"].notna()]
summary = dict(
    projects_construction=int(df["project"].nunique()), leaf_activities=int(len(df)), labeled=int(len(lab)),
    projects_with_labels=int(lab["project"].nunique()),
    by_group={g: dict(projects=int(x["project"].nunique()), activities=int(len(x)), labeled=int(x["is_late"].notna().sum()))
              for g, x in df.groupby("sector_group")},
    late_rate=round(float(lab["is_late"].mean()), 4), late7_rate=round(float(lab["is_late_7"].mean()), 4),
    start_late_rate=round(float(lab["start_late"].mean()), 4), grew_rate=round(float(lab["grew"].mean()), 4),
    own_delay_rate=round(float(lab["own_delay"].mean()), 4),
    late_given_start_late=round(float(lab.loc[lab["start_late"] == 1, "is_late"].mean()), 4),
    late_given_start_on_time=round(float(lab.loc[lab["start_late"] == 0, "is_late"].mean()), 4),
    median_delay_days=float(lab["delay_days"].median()), median_start_delay_days=float(lab["start_delay_days"].median()),
    float_nan_share=round(float(df["total_float_hr"].isna().mean()), 4), projects_with_network_cycle=cycles, baseline_end_derived_cells=derived_end[0],
    snapshots_rows=int(len(sn)), tracking_periods=int(sn.groupby(["project", "tp"]).ngroups) if len(sn) else 0,
    activities_with_predecessors_share=round(float((df["n_pred"] > 0).mean()), 4))
json.dump(summary, open(OUT / "dslib_etl_summary.json", "w"), indent=1, default=str)
pd.DataFrame(proj_summ).to_csv(OUT / "dslib_projects.csv", index=False)
print(json.dumps(summary, indent=1, default=str))
